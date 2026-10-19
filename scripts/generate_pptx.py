"""
Generate Globathon_Advisor_Dynamic_Summary.pptx for Advisor Pulse.
Uses live API demo metrics when available; falls back to recomputing from data.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "Globathon_Advisor_Dynamic_Summary.pptx"
SHOTS = ROOT / "deliverables" / "prototype_screenshots"
sys.path.insert(0, str(ROOT / "backend"))

NAVY = RGBColor(0x0B, 0x12, 0x20)
PANEL = RGBColor(0x16, 0x20, 0x33)
TEAL = RGBColor(0x2D, 0xD4, 0xBF)
WHITE = RGBColor(0xE8, 0xEE, 0xF8)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
ACCENT2 = RGBColor(0x3B, 0x82, 0xF6)
CRIT = RGBColor(0xF8, 0x71, 0x71)
OK = RGBColor(0x34, 0xD3, 0x99)


def set_slide_bg(slide, color=NAVY):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, left, top, width, height, text, size=18, bold=False, color=WHITE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Calibri"
    p.alignment = align
    return box


def add_footer(slide, text="Advisor Pulse · Globathon 2026 · DEMO data · Citations on-slide"):
    add_text(slide, Inches(0.4), Inches(7.05), Inches(12.5), Inches(0.3), text, size=10, color=MUTED)


def add_cite(slide, text, top=Inches(6.7)):
    add_text(slide, Inches(0.4), top, Inches(12.5), Inches(0.35), text, size=9, color=MUTED)


def rect(slide, left, top, width, height, fill=PANEL):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def load_metrics():
    try:
        import urllib.request

        with urllib.request.urlopen("http://127.0.0.1:8000/api/overview", timeout=3) as r:
            overview = json.loads(r.read().decode())
        with urllib.request.urlopen("http://127.0.0.1:8000/api/investors/DEMO-INV-001", timeout=3) as r:
            priya = json.loads(r.read().decode())
        with urllib.request.urlopen("http://127.0.0.1:8000/api/investors/DEMO-INV-002", timeout=3) as r:
            arjun = json.loads(r.read().decode())
        return overview, priya, arjun
    except Exception:
        from analytics.portfolio import DataStore, investor_snapshot
        from prioritization.engine import overview_kpis, rank_book, score_investor
        from summaries.generator import generate

        store = DataStore.load()
        ranked = rank_book(store)
        overview = overview_kpis(ranked)
        snap = investor_snapshot(store, "DEMO-INV-001")
        pri = score_investor(store, "DEMO-INV-001")
        priya = {"snapshot": snap, "priority": pri, "summary": generate(snap, pri)}
        snap2 = investor_snapshot(store, "DEMO-INV-002")
        pri2 = score_investor(store, "DEMO-INV-002")
        arjun = {"snapshot": snap2, "priority": pri2, "summary": generate(snap2, pri2)}
        return overview, priya, arjun


def try_add_image(slide, path: Path, left, top, width):
    if path.exists():
        slide.shapes.add_picture(str(path), left, top, width=width)
        return True
    return False


def build():
    overview, priya, arjun = load_metrics()
    dist = overview.get("priority_distribution", {})
    pr = priya["priority"]
    ps = priya["snapshot"]
    gp = ps["goal_progress"]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # ---- Slide 1 Title ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    rect(s, Inches(0), Inches(0), Inches(0.18), Inches(7.5), TEAL)
    add_text(s, Inches(0.6), Inches(1.6), Inches(8), Inches(0.5), "ADVISOR PULSE", size=18, bold=True, color=TEAL)
    add_text(s, Inches(0.6), Inches(2.1), Inches(10), Inches(1.2), "Where does this client stand?", size=40, bold=True)
    add_text(
        s,
        Inches(0.6),
        Inches(3.4),
        Inches(10),
        Inches(0.8),
        "AI-powered investor intelligence & action-prioritization layer for advisors — detect, explain, prioritize, prepare, act.",
        size=18,
        color=MUTED,
    )
    add_text(
        s,
        Inches(0.6),
        Inches(4.5),
        Inches(10),
        Inches(0.6),
        f"Demo book: {overview['total_clients']} clients · {overview['clients_requiring_review']} need review · "
        f"Critical {dist.get('Critical', 0)} · High {dist.get('High', 0)}",
        size=16,
        color=WHITE,
    )
    add_text(
        s,
        Inches(0.6),
        Inches(5.4),
        Inches(10),
        Inches(0.5),
        "Team: [Name 1]  ·  [Name 2]  ·  [Name 3]  ·  [Name 4]   |   Globathon 2026 FinTech",
        size=14,
        color=MUTED,
    )
    add_footer(s)

    # ---- Slide 2 Problem ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    add_text(s, Inches(0.5), Inches(0.35), Inches(12), Inches(0.6), "The problem: data exists — attention doesn't scale", size=28, bold=True)
    bullets = [
        "Holdings, goals, past advice, and notes live in different systems.",
        "Advisors can know their top ~20 clients deeply — the other ~200 rarely get proactive review.",
        "Hypothesis: the gap is not another chart — it is continuous, explainable prioritization of whom to contact next.",
        "Cerulli (2026): 64% of surveyed RIAs say AI reduced admin work; firms use AI to expand capacity, not replace advice.",
    ]
    y = 1.3
    for b in bullets:
        rect(s, Inches(0.5), Inches(y), Inches(12.2), Inches(0.95), PANEL)
        add_text(s, Inches(0.7), Inches(y + 0.25), Inches(11.8), Inches(0.6), "•  " + b, size=16)
        y += 1.1
    add_cite(s, "Sources: Globathon problem framing; Cerulli Associates press release, 9 Sep 2026 (AI capacity survey, n=68 RIAs). Client-book 20/200 = illustrative.")
    add_footer(s)

    # ---- Slide 3 Existing workflows ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    add_text(s, Inches(0.5), Inches(0.35), Inches(12), Inches(0.6), "Why CRM + portfolio tools still leave a gap", size=28, bold=True)
    cols = [
        ("Store & engage", "CRM (e.g. Salesforce FSC)\nHousehold 360, tasks, AI meeting prep"),
        ("Analyze & construct", "Morningstar / Aladdin Wealth\nResearch, risk, proposals, rebalancing"),
        ("Still missing", "Cross-context attention loop\nGoal path vs actual + explainable why-now + long-tail queue"),
    ]
    for i, (h, body) in enumerate(cols):
        left = 0.5 + i * 4.2
        rect(s, Inches(left), Inches(1.3), Inches(3.9), Inches(4.2), PANEL)
        add_text(s, Inches(left + 0.25), Inches(1.55), Inches(3.4), Inches(0.5), h, size=18, bold=True, color=TEAL)
        add_text(s, Inches(left + 0.25), Inches(2.3), Inches(3.4), Inches(2.8), body, size=15, color=WHITE)
    add_cite(s, "Sources: salesforce.com/financial-services/wealth-management-software; morningstar Advisor Workstation; blackrock.com/aladdin (Aladdin Wealth). Fair comparison — not capability denial.")
    add_footer(s)

    # ---- Slide 4 Solution ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    add_text(s, Inches(0.5), Inches(0.35), Inches(12), Inches(0.5), "Advisor Pulse — the attention layer", size=28, bold=True)
    steps = ["Detect", "Explain", "Prioritize", "Prepare", "Act", "Learn"]
    for i, step in enumerate(steps):
        left = 0.5 + i * 2.1
        rect(s, Inches(left), Inches(1.3), Inches(1.9), Inches(0.7), TEAL if i % 2 == 0 else ACCENT2)
        add_text(s, Inches(left), Inches(1.4), Inches(1.9), Inches(0.5), step, size=16, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    caps = [
        "Unified client context: holdings + goals + advice + notes (+ optional cash-flow signals)",
        "Expectation-vs-actual with evidence-linked driver analysis (qualified, not fabricated)",
        "Transparent multi-factor priority (Critical / High / Monitor / On Track) — configurable thresholds",
        "Meeting briefs & review suggestions — human approval required; no autonomous trading",
    ]
    y = 2.4
    for c in caps:
        add_text(s, Inches(0.6), Inches(y), Inches(12), Inches(0.45), "→  " + c, size=16)
        y += 0.7
    add_footer(s)

    # ---- Slide 5 Architecture ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    add_text(s, Inches(0.5), Inches(0.35), Inches(12), Inches(0.5), "How the intelligence works", size=28, bold=True)
    flow = [
        "Txn + portfolio sources",
        "Validate",
        "Goal & portfolio analytics",
        "Attention model",
        "Evidence-linked summary",
        "Advisor approves",
    ]
    for i, f in enumerate(flow):
        top = 1.4 + (i % 3) * 1.5
        left = 0.6 if i < 3 else 7.0
        rect(s, Inches(left), Inches(top), Inches(5.5), Inches(1.2), PANEL)
        add_text(s, Inches(left + 0.3), Inches(top + 0.35), Inches(5), Inches(0.5), f"{i+1}. {f}", size=16, bold=True)
    add_cite(s, "Security boundary: no sensitive investor data to external LLMs without approved DPA. Demo uses deterministic templates. SEBI IA Reg. 15(14): IA remains responsible for AI tools.")
    add_footer(s)

    # ---- Slide 6 Value 20 vs 200 ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    add_text(s, Inches(0.5), Inches(0.35), Inches(12), Inches(0.5), "Where the value lies: 20 vs 200", size=28, bold=True)
    rect(s, Inches(0.5), Inches(1.2), Inches(6), Inches(4.5), PANEL)
    add_text(s, Inches(0.75), Inches(1.4), Inches(5.5), Inches(0.4), "Known clients (illustrative top 20)", size=18, bold=True, color=TEAL)
    add_text(
        s,
        Inches(0.75),
        Inches(2.0),
        Inches(5.5),
        Inches(3.2),
        "Meeting prep briefs\nChanges since last review\nGoal progress + evidence\nConversation-ready context\n\nSupports depth — does not replace relationship.",
        size=15,
    )
    rect(s, Inches(6.8), Inches(1.2), Inches(6), Inches(4.5), PANEL)
    add_text(s, Inches(7.05), Inches(1.4), Inches(5.5), Inches(0.4), "Long tail (illustrative ~200)", size=18, bold=True, color=TEAL)
    add_text(
        s,
        Inches(7.05),
        Inches(2.0),
        Inches(5.5),
        Inches(3.2),
        f"Proactive detection of material deviations\nDemo: {overview['long_tail_count']} long-tail · {overview['clients_requiring_review']} total need review\nConfigurable thresholds reduce alert fatigue\nData-quality ≠ financial distress\n\nGreatest incremental value: scaling quality attention.",
        size=15,
    )
    add_cite(s, "Illustrative book sizes. Industry context: Institutional Investor summarizing Cerulli — senior advisors ~160 clients; 29% manage >200. Not a measured productivity claim for Advisor Pulse.")
    add_footer(s)

    # ---- Slide 7 Investor summary ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    add_text(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.45), "Investor Dynamic Summary — Priya Sharma (DEMO)", size=26, bold=True)
    img = SHOTS / "investor_summary.png"
    if not try_add_image(s, img, Inches(0.4), Inches(0.9), Inches(7.8)):
        rect(s, Inches(0.4), Inches(0.9), Inches(7.8), Inches(5.5), PANEL)
        add_text(s, Inches(0.7), Inches(3.2), Inches(7), Inches(0.5), "[Screenshot placeholder — run prototype]", size=16, color=MUTED, align=PP_ALIGN.CENTER)
    rect(s, Inches(8.5), Inches(0.9), Inches(4.4), Inches(5.5), PANEL)
    add_text(s, Inches(8.7), Inches(1.1), Inches(4), Inches(0.4), "Computed demo metrics", size=16, bold=True, color=TEAL)
    metrics_txt = (
        f"Priority: {pr['priority']} (score {pr['score']})\n"
        f"Portfolio ≈ ₹{ps['portfolio_value']/1e5:.1f}L\n"
        f"Vs expected path: {gp['deviation_pct_vs_expected']:+.1f}%\n"
        f"Active return: {ps['performance'].get('active_return_pct')} pp\n"
        f"Goal: {gp['goal_name']}\n"
        f"Confidence: {ps['gap_explanation'].get('confidence')}\n\n"
        "SYNTHETIC DEMO — not a real client."
    )
    add_text(s, Inches(8.7), Inches(1.7), Inches(4), Inches(4.2), metrics_txt, size=14)
    add_footer(s)

    # ---- Slide 8 Live demo ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    add_text(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.45), "Live demo walkthrough", size=28, bold=True)
    story = [
        f"1. DETECT — Priya behind expected path by {gp['deviation_pct_vs_expected']:.1f}%",
        "2. EXPLAIN — Mid-cap overweight + active TWR lag vs benchmark (qualified drivers)",
        f"3. PRIORITIZE — Surfaces as {pr['priority']} in book of {overview['total_clients']}",
        "4. PREPARE — Meeting brief with open questions & overdue follow-ups",
        "5. ACT — Advisor review suggestions only (rebalance discussion, close follow-up)",
        "6. CONTRAST — Arjun Mehta: AUM dip from planned withdrawal → NOT escalated",
    ]
    y = 1.0
    for line in story:
        rect(s, Inches(0.5), Inches(y), Inches(12.2), Inches(0.8), PANEL)
        add_text(s, Inches(0.7), Inches(y + 0.2), Inches(11.8), Inches(0.5), line, size=15)
        y += 0.9
    add_footer(s)

    # ---- Slide 9 Privacy ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    add_text(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.45), "Privacy, trust & human oversight", size=28, bold=True)
    items = [
        "RBAC, encryption in transit/at rest, audit trails — production architecture (prototype is local demo)",
        "Data minimization & purpose limitation (DPDP Act 2023 design considerations)",
        "SEBI IA Reg. 15(14): IA solely responsible for AI tools & client data integrity; disclose AI use (18(9))",
        "No investor PII to external models without approved processing arrangement — demo uses templates",
        "Meeting transcripts: mocked local path in prototype; production = on-prem ASR + consent",
        "Prompt-injection resistant summaries: numbers only from structured analytics",
        "Honest limit: NOT claimed compliant, certified, or regulator-approved",
    ]
    y = 1.0
    for it in items:
        add_text(s, Inches(0.6), Inches(y), Inches(12), Inches(0.55), "•  " + it, size=14)
        y += 0.7
    add_cite(s, "Sources: SEBI IA Regulations PDF; SEBI IA Master Circular PDF; MeitY DPDP Act 2023 PDF. Design considerations ≠ compliance certification.")
    add_footer(s)

    # ---- Slide 10 Competitive ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    add_text(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.45), "Competitive positioning", size=28, bold=True)
    add_text(
        s,
        Inches(0.5),
        Inches(1.0),
        Inches(12.2),
        Inches(5.2),
        "Capability                    Salesforce FSC/Agentforce   Morningstar AWS   Aladdin Wealth   Advisor Pulse\n"
        "CRM / household 360           Y                           P                 P                Integrates\n"
        "Deep risk / stress            NV                          P                 Y                Not claimed\n"
        "Goal path vs expected         P                           P                 P                Core\n"
        "Explainable gap drivers       NV                          P                 P                Core (qualified)\n"
        "Transparent priority factors  NV                          NV                P (alerts)       Core\n"
        "Evidence-bound NL briefs      Y                           NV                NV               Y (structured→text)\n"
        "Long-tail attention workflow  P                           NV                P                Primary job-to-be-done\n\n"
        "Differentiation: combination + explainable attention loop — not 'another dashboard'.\n"
        "Y=evidenced · P=partial · NV=not verified in sources examined. See research/competitor_matrix.md.",
        size=13,
    )
    add_footer(s)

    # ---- Slide 11 Roadmap ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    add_text(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.45), "Implementation & validation roadmap", size=28, bold=True)
    rect(s, Inches(0.5), Inches(1.1), Inches(4), Inches(5), PANEL)
    add_text(s, Inches(0.7), Inches(1.3), Inches(3.6), Inches(0.4), "Built today", size=16, bold=True, color=TEAL)
    add_text(
        s,
        Inches(0.7),
        Inches(1.9),
        Inches(3.6),
        Inches(3.8),
        "• React + FastAPI prototype\n• Synthetic DEMO book (220)\n• TWR + Modified Dietz\n• Transparent prioritization\n• Meeting briefs + notes\n• Deterministic summaries\n• Tests + research pack",
        size=14,
    )
    rect(s, Inches(4.7), Inches(1.1), Inches(4), Inches(5), PANEL)
    add_text(s, Inches(4.9), Inches(1.3), Inches(3.6), Inches(0.4), "Needs integration", size=16, bold=True, color=TEAL)
    add_text(
        s,
        Inches(4.9),
        Inches(1.9),
        Inches(3.6),
        Inches(3.8),
        "• Live custodian / PMS feeds\n• CRM task push\n• On-prem ASR\n• Approved LLM under DPA\n• Full Brinson attribution\n• Production security review",
        size=14,
    )
    rect(s, Inches(8.9), Inches(1.1), Inches(4), Inches(5), PANEL)
    add_text(s, Inches(9.1), Inches(1.3), Inches(3.6), Inches(0.4), "Pilot metrics (plan)", size=16, bold=True, color=TEAL)
    add_text(
        s,
        Inches(9.1),
        Inches(1.9),
        Inches(3.6),
        Inches(3.8),
        "• Review prep time\n• % alerts advisor-validated\n• Missed events (retro)\n• False-positive rate\n• % summaries with cited facts\n\nNo invented pilot results.",
        size=14,
    )
    add_cite(s, "Round 2 Globathon rule: problem-specific logic must be built during the 12-hour finale; this repo is Round 1 ideation + reference architecture.")
    add_footer(s)

    # ---- Slide 12 Close ----
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    add_text(s, Inches(0.6), Inches(2.0), Inches(12), Inches(0.6), "Where does this client stand?", size=36, bold=True, color=TEAL)
    add_text(
        s,
        Inches(0.6),
        Inches(2.9),
        Inches(12),
        Inches(1.2),
        "Advisor Pulse helps advisors spend less time searching for signals —\nand more time acting on the right ones.",
        size=22,
    )
    add_text(
        s,
        Inches(0.6),
        Inches(4.5),
        Inches(12),
        Inches(0.8),
        "Detect → Explain → Prioritize → Prepare → Act → Learn\nTeam placeholders · Globathon 2026 · React web prototype",
        size=16,
        color=MUTED,
    )
    add_footer(s)

    # Optional appendix slide — Round 2 short pitch note
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(s)
    add_text(s, Inches(0.5), Inches(0.3), Inches(12), Inches(0.45), "Appendix — Finale pitch cut (2–4 slides)", size=26, bold=True)
    add_text(
        s,
        Inches(0.6),
        Inches(1.2),
        Inches(12),
        Inches(5),
        "Official Round 2 pitch requires a short PPT (2–4 slides). Suggested cut:\n\n"
        "1) Title + promise (Slide 1)\n"
        "2) Solution workflow + differentiation (Slide 4)\n"
        "3) Demo metrics / Priya story (Slides 7–8 compressed)\n"
        "4) Trust + ask (Slide 9 + close)\n\n"
        "Also required at finale: working prototype, source code, run instructions, live demo + Q&A.\n"
        "Reuse rule: problem-specific features must be developed during the coding round.",
        size=16,
    )
    add_footer(s)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
